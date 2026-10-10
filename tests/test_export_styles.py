"""C01: installed ZIPs, stable layouts and editable Word structure."""
import tempfile
import unittest
import zipfile
from pathlib import Path
from mdreader import core
from tools.verify_delivery import ROOT, install_exports, inspect_pdf


class ExportStyleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='export-styles-')
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.ws = core.Workspace(str(self.folder / 'workspace'))
        self.addCleanup(self.ws.plugins.shutdown)
        install_exports(self.ws)
        self.api = core.Api(self.ws, str(ROOT/'webui'))
        self.doc = self.folder/'sample.md'
        self.doc.write_text('# Report\n\n正文 123。\n\n## Section\n\n| name | value |\n|---|---|\n| item | 42 |\n\n```python\nprint(42)\n```\n', encoding='utf-8')
        self.api.loose.open_path(str(self.doc))

    def export(self, style, ext, suffix=''):
        target = self.folder/('%s%s.%s' % (style, suffix, ext))
        command = next(row['command'] for row in self.ws.plugins.commands() if row['extension'] == ext)
        self.api.post('/api/plugins/export', {'command': command, 'doc': str(self.doc),
            'dest': str(target), 'options': {'style': style}, 'wait': 120})
        return target

    def test_word_styles_are_distinct_and_remain_editable_after_reopen(self):
        from docx import Document
        sizes = []
        for style in ('report', 'technical'):
            doc = Document(self.export(style, 'docx'))
            sizes.append(doc.styles['Heading 1'].font.size.pt)
            expected = ('Report', 'Section', 'item', '42')
            def content(document):
                return '\n'.join(p.text for p in document.paragraphs) + '\n'.join(
                    cell.text for table in document.tables for row in table.rows for cell in row.cells)
            for text in expected:
                self.assertIn(text, content(doc))
            saved = self.folder/(style+'-reopened.docx')
            doc.save(saved)
            self.assertEqual(content(doc), content(Document(saved)))
            self.assertEqual(len(Document(saved).tables), 1)
            self.assertGreater(doc.styles['Heading 1'].font.size.pt, doc.styles['Heading 2'].font.size.pt)
        self.assertNotEqual(*sizes)

    def test_repeat_pdf_layout_does_not_follow_reading_theme(self):
        for style in ('report', 'technical'):
            core.write_ui_settings(self.ws.root, 'light')
            first = inspect_pdf(self.export(style, 'pdf'), render=False)
            core.write_ui_settings(self.ws.root, 'dark')
            second = inspect_pdf(self.export(style, 'pdf', '-again'), render=False)
            self.assertEqual(first, second)
            self.assertFalse(first['content_violations'])

    def test_word_structure_and_styles_do_not_follow_reading_theme(self):
        def structure(path):
            with zipfile.ZipFile(path) as archive:
                return tuple(archive.read(name) for name in ('word/document.xml', 'word/styles.xml'))
        for style in ('report', 'technical'):
            core.write_ui_settings(self.ws.root, 'light')
            original = structure(self.export(style, 'docx'))
            for theme in ('dark', 'eye'):
                core.write_ui_settings(self.ws.root, theme)
                self.assertEqual(original, structure(self.export(style, 'docx', '-'+theme)))

    def test_invalid_style_does_not_touch_existing_target(self):
        target = self.folder/'unknown.pdf'
        target.write_bytes(b'original-output')
        with self.assertRaises(ValueError):
            self.export('unknown', 'pdf')
        self.assertEqual(target.read_bytes(), b'original-output')

    def test_compact_style_keeps_clearance_below_wide_formulas(self):
        from tools.verify_delivery import FIXTURES
        self.doc.write_text((FIXTURES/'24-wide-formula.md').read_text(encoding='utf-8'), encoding='utf-8')
        inspection = inspect_pdf(self.export('technical', 'pdf'), render=False)
        self.assertFalse(inspection['content_violations'])
        self.assertFalse(inspection['image_text_overlap'])
