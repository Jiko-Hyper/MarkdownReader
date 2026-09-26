"""End-to-end tests against the shipped ZIP and real worker processes."""
import os
from pathlib import Path
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from mdreader import core, plugins
from mdreader.conversion import convert_file

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "plugins/dependencies"))


class ConversionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.temp.name)
        cls.ws = core.Workspace(str(cls.root / "workspace"))
        cls.ws.plugins.install(str(ROOT / "plugins/packages/mdreader.document-convert-1.0.0.zip"))
        cls.ws.plugins.enable("mdreader.document-convert")
        cls.deps = cls.ws.plugins.entry("mdreader.document-convert")["path"]
        sys.path.insert(0, cls.deps)
        cls.commands = {c["format"]: c["command"] for c in cls.ws.plugins.commands()
                        if c["plugin"] == "mdreader.document-convert"}

    @classmethod
    def tearDownClass(cls):
        cls.ws.plugins.shutdown()
        sys.path.remove(cls.deps)
        cls.temp.cleanup()

    def setUp(self):
        self.folder = Path(tempfile.mkdtemp(dir=self.root))

    def convert(self, fmt, source, name=None):
        target = self.folder / (name or ("output.html" if fmt == "md-html" else "output.md"))
        result = convert_file(self.ws.plugins, self.commands[fmt], str(source), str(target))
        return target.read_text(encoding="utf-8"), result

    def image(self):
        from PIL import Image
        target = self.folder / "中文图片.png"
        Image.new("RGB", (10, 20), "blue").save(target)
        return target

    def test_markdown_html_roundtrip_structure_unicode_and_embedded_image(self):
        image = self.image()
        source = self.folder / "中文.md"
        source.write_text('# 标题\n\n**粗体** 和 *强调*\n\n- 一\n- 二\n\n'
                          '| 名称 | 数量 |\n|---|---|\n| 中文 | 42 |\n\n'
                          '```python\nprint("你好")\n```\n\n![图片](中文图片.png)\n', encoding="utf-8")
        original = source.read_bytes()
        output, result = self.convert("md-html", source)
        self.assertIn("<h1>标题</h1>", output)
        self.assertIn("<table>", output)
        self.assertIn("data:image/png;base64,", output)
        self.assertIn("Content-Security-Policy", output)
        self.assertEqual(source.read_bytes(), original)
        back, result = self.convert("html-md", Path(result["path"]), "back.md")
        for text in ("# 标题", "**粗体**", "42", 'print("你好")'):
            self.assertIn(text, back)
        self.assertEqual(next(Path(result["assets"]).iterdir()).read_bytes(), image.read_bytes())

    def test_word_headings_table_lists_image(self):
        from docx import Document
        doc = Document()
        doc.add_heading("中文标题", 1)
        doc.add_paragraph("第一项", style="List Bullet")
        doc.add_paragraph().add_run("粗体内容").bold = True
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "品名"
        table.cell(0, 1).text = "数量"
        table.cell(1, 0).text = "苹果"
        table.cell(1, 1).text = "42"
        image = self.image()
        doc.add_picture(str(image))
        source = self.folder / "测试.docx"
        doc.save(source)
        text, result = self.convert("docx-md", source)
        for expected in ("# 中文标题", "- 第一项", "**粗体内容**", "苹果", "42", "image-001.png"):
            self.assertIn(expected, text)
        self.assertEqual(next(Path(result["assets"]).iterdir()).read_bytes(), image.read_bytes())

    def test_pdf_real_text_pages(self):
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        source = self.folder / "source.pdf"
        canvas = Canvas(str(source))
        canvas.drawString(30, 700, "Hello PDF page one")
        canvas.showPage()
        canvas.drawString(30, 700, "Second page 42")
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        canvas.setFont("STSong-Light", 12)
        canvas.drawString(30, 670, "中文内容")
        canvas.save()
        text, result = self.convert("pdf-md", source)
        self.assertIn("Hello PDF page one", text)
        self.assertIn("Second page 42", text)
        self.assertIn("中文内容", text)
        self.assertIn("第 2 页", text)
        self.assertTrue(result["warnings"])

    def test_scanned_pdf_encrypted_pdf_and_bad_docx_fail_without_output(self):
        from pypdf import PdfWriter
        for name, encrypted in (("blank.pdf", False), ("encrypted.pdf", True)):
            source = self.folder / name
            writer = PdfWriter()
            writer.add_blank_page(width=200, height=200)
            if encrypted:
                writer.encrypt("secret")
            writer.write(source)
            with self.assertRaises(plugins.TaskError):
                self.convert("pdf-md", source)
        bad = self.folder / "bad.docx"
        bad.write_bytes(b"not a word document")
        with self.assertRaises(plugins.TaskError):
            self.convert("docx-md", bad)
        self.assertFalse((self.folder / "output.md").exists())

    def test_html_sanitizes_active_content_and_does_not_read_outside_root(self):
        outside = self.root / "private.png"
        outside.write_bytes(self.image().read_bytes())
        source = self.folder / "source.html"
        source.write_text('<h1>正文</h1><script>secret_script()</script>'
                          '<a href="javascript:alert(1)">链接</a><img src="../private.png" alt="私有">'
                          '<img src="https://example.invalid/image.png" alt="远程">', encoding="utf-8")
        text, result = self.convert("html-md", source)
        self.assertIn("# 正文", text)
        self.assertNotIn("secret_script", text)
        self.assertNotIn("javascript:", text)
        self.assertFalse(result["assets"])
        self.assertTrue(result["warnings"])

    def test_existing_target_and_racing_target_are_not_overwritten(self):
        source = self.folder / "source.html"
        source.write_text("<p>Hello</p>", encoding="utf-8")
        target = self.folder / "output.md"
        target.write_text("keep me", encoding="utf-8")
        with self.assertRaises(FileExistsError):
            self.convert("html-md", source)
        self.assertEqual(target.read_text(), "keep me")
        target.unlink()
        original_link = os.link
        def competing_link(temp, dest):
            Path(dest).write_text("another process", encoding="utf-8")
            return original_link(temp, dest)
        with patch("mdreader.conversion.os.link", side_effect=competing_link):
            with self.assertRaises(FileExistsError):
                self.convert("html-md", source)
        self.assertEqual(target.read_text(), "another process")
        self.assertFalse(list(self.folder.glob(".conversion-*")))

    def test_publish_failure_cleans_images_and_temp_files(self):
        self.image()
        source = self.folder / "source.html"
        source.write_text('<p>正文</p><img src="中文图片.png">', encoding="utf-8")
        with patch("mdreader.conversion.os.link", side_effect=OSError("disk failure")):
            with self.assertRaises(OSError):
                self.convert("html-md", source)
        self.assertFalse(list(self.folder.glob("conversion-assets-*")))
        self.assertFalse(list(self.folder.glob(".conversion-*")))
        self.assertFalse((self.folder / "output.md").exists())

    def test_empty_and_gb18030_html(self):
        source = self.folder / "source.html"
        source.write_bytes("<h1>中文旧编码</h1>".encode("gb18030"))
        text, _ = self.convert("html-md", source)
        self.assertIn("中文旧编码", text)
        source.write_text("", encoding="utf-8")
        text, result = self.convert("html-md", source, "empty.md")
        self.assertFalse(text.strip())
        self.assertTrue(result["warnings"])

    def test_disabled_plugin_rejects_conversion(self):
        source = self.folder / "source.md"
        source.write_text("# hello", encoding="utf-8")
        self.ws.plugins.disable("mdreader.document-convert")
        try:
            with self.assertRaises(plugins.PluginError):
                self.convert("md-html", source)
        finally:
            self.ws.plugins.enable("mdreader.document-convert")

    def test_desktop_menu_dispatch_and_disable(self):
        from mdreader.winui import MarkdownWindow
        owner = SimpleNamespace(ws=self.ws, manage_plugins=lambda: None,
                                insert_image_with_plugin=lambda: None)
        owner.plugin_commands = lambda cap: MarkdownWindow.plugin_commands(owner, cap)
        with patch("mdreader.conversion.show_conversion") as show:
            entries = MarkdownWindow.plugin_menu_entries(owner)
            self.assertEqual(len(entries["conversions"]), 4)
            for label, action in entries["conversions"]:
                action()
                self.assertEqual(show.call_args.args[1]["title"], label)
        self.ws.plugins.disable("mdreader.document-convert")
        try:
            self.assertEqual(MarkdownWindow.plugin_menu_entries(owner)["conversions"], [])
        finally:
            self.ws.plugins.enable("mdreader.document-convert")

    def test_desktop_dialog_converts_and_cancel_does_not_write(self):
        from mdreader.conversion import show_conversion
        from unittest.mock import Mock
        source = self.folder / "source.md"
        source.write_text("# 菜单转换", encoding="utf-8")
        target = self.folder / "dialog.html"
        owner = SimpleNamespace(root=None, ws=self.ws, notice=Mock())
        command = self.ws.plugins.find_command(self.commands["md-html"])
        with patch("tkinter.filedialog.askopenfilename", return_value=str(source)), \
             patch("tkinter.filedialog.asksaveasfilename", return_value=str(target)), \
             patch("mdreader.media_ui.run_job", side_effect=lambda owner, fn: fn()), \
             patch("tkinter.messagebox.showinfo") as info, \
             patch("tkinter.messagebox.showerror") as error:
            show_conversion(owner, command)
            self.assertIn("<h1>菜单转换</h1>", target.read_text(encoding="utf-8"))
            info.assert_called_once()
            error.assert_not_called()
        with patch("tkinter.filedialog.askopenfilename", return_value=""), \
             patch("mdreader.conversion.convert_file") as convert:
            show_conversion(owner, command)
            convert.assert_not_called()


if __name__ == "__main__":
    unittest.main()
