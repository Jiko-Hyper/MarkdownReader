"""Frozen 20-document acceptance corpus; expected content precedes conversion.

No OCR/legacy DOC inputs are counted as successful conversion samples.
Every document must pass every applicable assertion; no average hides a failure.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
import sys
import tempfile
import time
import unicodedata

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / "plugins/dependencies"), str(ROOT / "plugins/conversion-dependencies")]
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from markdown_it import MarkdownIt
from mdreader import core
from mdreader.conversion import convert_file


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def norm(text):
    # Only Unicode canonical equivalence and whitespace are ignored, not punctuation or values.
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", text))


def make_corpus(folder):
    folder.mkdir(parents=True, exist_ok=True)
    pictures = []
    for n, color in enumerate(("#2874a6", "#e67e22"), 1):
        im = Image.new("RGB", (200, 100), "white")
        draw = ImageDraw.Draw(im)
        draw.rectangle((10, 10, 100 + n * 25, 80), fill=color)
        draw.text((15, 35), "Sample %d" % n, fill="black")
        path = folder / ("sample-%d.png" % n)
        im.save(path)
        pictures.append(path)
    cases = []
    def case(index, name, level, ext):
        row = dict(id="%02d" % index, name=name, complexity=level, format=ext,
                   source="%02d-%s.%s" % (index, name, ext), texts=[], selectors={},
                   table_rows=[], links=[], codes=[], pictures=[], order=True)
        cases.append(row)
        return row
    def para(doc, row, text, style=None):
        row["texts"].append(text)
        return doc.add_paragraph(text, style)
    def head(doc, row, text, level=1):
        row["texts"].append(text)
        doc.add_heading(text, level)
        row["selectors"]["h%d" % level] = row["selectors"].get("h%d" % level, 0) + 1
    def table(doc, row, values):
        t = doc.add_table(rows=len(values), cols=len(values[0]))
        t.style = "Table Grid"
        for cells, source in zip(t.rows, values):
            for cell, text in zip(cells.cells, source):
                cell.text = text
                row["texts"].append(text)
        row["table_rows"].extend(values)
        row["selectors"]["table"] = row["selectors"].get("table", 0) + 1
        return t
    def pic(doc, row, n):
        doc.add_picture(str(pictures[n]), width=Inches(2))
        row["pictures"].append(pictures[n].name)

    for index in range(1, 11):
        names = ["日常通知", "分层方案", "样式与链接", "多级任务列表", "库存明细表",
                 "代码与引用", "双图片说明", "合并单元格", "多语言数值", "长篇混合报告"]
        row = case(index, names[index-1], "简单" if index < 3 else "中等" if index < 7 else "复杂", "docx")
        doc = Document()
        head(doc, row, names[index-1])
        if index == 1:
            para(doc, row, "会议时间为周五下午三点，地点为研发中心。")
            para(doc, row, "Please bring the revised schedule and meeting notes.")
        elif index == 2:
            for n in range(1, 5):
                head(doc, row, "阶段 %d" % n, min(n+1, 4))
                para(doc, row, "本阶段验收编号 STAGE-%02d。" % n)
        elif index == 3:
            p = doc.add_paragraph()
            for text, attr in [("重点条款", "bold"), ("补充说明", "italic")]:
                setattr(p.add_run(text), attr, True)
                row["texts"].append(text)
            row["selectors"].update(strong=1, em=1)
            p = doc.add_paragraph()
            link = OxmlElement("w:hyperlink")
            rel = doc.part.relate_to("https://example.org/spec?q=1&lang=zh",
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link.set(qn("r:id"), rel)
            run, text = OxmlElement("w:r"), OxmlElement("w:t")
            text.text = "规范链接"
            run.append(text); link.append(run); p._p.append(link)
            row["texts"].append("规范链接")
            row["links"].append("https://example.org/spec?q=1&lang=zh")
        elif index == 4:
            for text in ("准备材料", "验证输入", "记录结果"):
                para(doc, row, text, "List Number")
            for text in ("附加条件甲", "附加条件乙"):
                para(doc, row, text, "List Bullet")
            row["selectors"].update(ol=1, ul=1, li=5)
        elif index == 5:
            table(doc, row, [["编码", "品名", "数量", "金额"]] +
                [["SKU-%02d" % n, "零件%d" % n, str(n*7), "%.2f" % (n*13.17)] for n in range(1, 13)])
        elif index == 6:
            doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
            code = 'def total(values):\n    return sum(x for x in values if x > 0)'
            para(doc, row, code, "Code")
            row["codes"].append(code)
            para(doc, row, "引用说明：输入必须为数值。", "Quote")
            row["selectors"]["blockquote"] = 1
        elif index == 7:
            for n in range(2):
                para(doc, row, "示意图 %d 展示了不同区域。" % (n+1))
                pic(doc, row, n)
        elif index == 8:
            # Explicit expected rectangular representation of a merged header.
            t = doc.add_table(rows=3, cols=3)
            t.style = "Table Grid"
            t.cell(0, 0).merge(t.cell(0, 1)).text = "合并区域"
            t.cell(0, 2).text = "指标"
            for r, values in enumerate([["东区", "A组", "17"], ["西区", "B组", "29"]], 1):
                for cell, text in zip(t.rows[r].cells, values): cell.text = text
            row["texts"] += ["合并区域", "指标", "东区", "A组", "17", "西区", "B组", "29"]
            row["table_rows"] = [["合并区域", "", "指标"], ["东区", "A组", "17"], ["西区", "B组", "29"]]
            row["selectors"]["table"] = 1
        elif index == 9:
            for text in ["中文、日本語、한국어、English、café。", "数值：−12.50，±0.25，42%，H₂O，x²。",
                         "语法字面量：a_b *literal* [x] #tag `tick` <value> | pipe。"]:
                para(doc, row, text)
        else:
            for section in range(1, 7):
                if section > 1: doc.add_page_break()
                head(doc, row, "第%d章 项目复盘" % section, 2)
                for n in range(1, 9):
                    para(doc, row, "章节%d段落%d：已完成设计、评审与验收，编号 R%02d-%02d。" % (section,n,section,n))
                table(doc, row, [["任务", "状态"], ["阶段%d" % section, "已完成"]])
                if section in (2, 5): pic(doc, row, section % 2)
        doc.save(folder / row["source"])

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdfmetrics.registerFont(TTFont("CorpusArial", "C:/Windows/Fonts/segoeui.ttf"))
    for index in range(11, 21):
        names = ["英文短文", "中文说明", "标题层次", "执行清单", "网格财务表",
                 "双图说明页", "双栏交错内容流", "六页长文", "代码和特殊字符", "图文表链接混排"]
        row = case(index, names[index-11], "简单" if index < 13 else "中等" if index < 17 else "复杂", "pdf")
        c = canvas.Canvas(str(folder / row["source"]), pagesize=(595, 842), invariant=1)
        y = 790
        def line(text, size=12, font="STSong-Light", x=45, at=None, heading=0):
            nonlocal y
            if at is not None: y = at
            c.setFont(font, size); c.drawString(x, y, text)
            row["texts"].append(text)
            if heading: row["selectors"]["h%d" % heading] = row["selectors"].get("h%d" % heading, 0) + 1
            y -= size + 12
        def ptable(values):
            nonlocal y
            top = y
            widths = [145] * len(values[0])
            c.setLineWidth(.6)
            for r, values_row in enumerate(values):
                xx = 45
                for text, width in zip(values_row, widths):
                    c.rect(xx, top - (r+1)*27, width, 27)
                    c.setFont("STSong-Light", 11)
                    c.drawString(xx+6, top-r*27-18, text)
                    row["texts"].append(text)
                    xx += width
            row["table_rows"].extend(values)
            row["selectors"]["table"] = row["selectors"].get("table", 0)+1
            y = top-len(values)*27-25
        def pimage(n):
            nonlocal y
            c.drawImage(str(pictures[n]), 45, y-100, width=200, height=100)
            row["pictures"].append(pictures[n].name)
            y -= 125
        if index == 11:
            line("Weekly operations memo", 22, "Helvetica", heading=1)
            line("Revenue increased by 12.5 percent in September.", font="Helvetica")
            line("Next review: Monday at 09:30.", font="Helvetica")
        elif index == 12:
            line("中文操作说明", 22, heading=1)
            for text in ("步骤一：检查设备状态和电源。", "步骤二：记录编号 CN-2026-0930。", "步骤三：保存结果并结束操作。"):
                line(text)
        elif index == 13:
            line("工程方案", 22, heading=1)
            for text in ("设计目标", "实施计划", "验收标准"):
                line(text, 17, heading=2); line(text + "包含负责人、截止日期和检查记录。")
        elif index == 14:
            line("执行清单", 22, heading=1)
            for text in ("1. 检查温度", "2. 读取压力", "3. 填写日志", "- 材料齐全", "- 仪器已校准"):
                line(text)
            row["texts"] = [re.sub(r"^(?:\d+\. |\- )", "", t) for t in row["texts"]]
            row["selectors"].update(ol=1, ul=1, li=5)
        elif index == 15:
            line("季度财务表", 22, heading=1)
            ptable([["项目", "收入", "成本"]] + [["部门%d" % n, "%.2f" % (n*1243.71), "%.2f" % (n*521.09)] for n in range(1, 9)])
        elif index == 16:
            line("双图说明", 22, heading=1)
            for n in range(2): line("图%d：区域分析。" % (n+1)); pimage(n)
        elif index == 17:
            line("双栏阅读顺序", 22, heading=1)
            # Deliberately interleave PDF drawing order, but expected reading order is columns.
            left, right = [], []
            for n in range(1, 13):
                l, r = "左栏第%02d段：分析结果。" % n, "右栏第%02d段：后续计划。" % n
                line(l, x=45, at=730-n*28); line(r, x=330, at=730-n*28)
                left.append(l); right.append(r)
            row["texts"] = [row["texts"][0]] + left + right
        elif index == 18:
            for page in range(1, 7):
                if page > 1: c.showPage(); y = 790
                line("第%d页 运行记录" % page, 22, heading=1)
                for n in range(1, 17): line("页%d行%02d：测量值 %d.%02d，状态正常。" % (page,n,page*17,n))
        elif index == 19:
            line("Technical examples", 22, "Helvetica", heading=1)
            code = 'def net(values):\n    return sum(values) - 12.50'
            for text in code.splitlines(): line(text, font="Courier")
            row["codes"].append(code)
            line("Greek: alpha α beta β; value −12.50 ± 0.25; H₂O x²", font="CorpusArial")
            line("Literal: a_b *star* [item] #tag `tick` <value> | pipe.", font="Helvetica")
        else:
            line("综合项目报告", 22, heading=1)
            line("项目结论", 17, heading=2)
            line("设备运行稳定，三个批次完成检查。")
            ptable([["批次", "合格数", "日期"], ["A-01", "128", "2026-09-01"], ["B-02", "256", "2026-09-02"]])
            line("图一展示了主要区域。"); pimage(0)
            pos = y
            line("访问项目规范")
            c.linkURL("https://example.org/project", (45,pos-3,129,pos+13), relative=0)
            row["links"].append("https://example.org/project")
            line("最终检查", 17, heading=2)
            line("1. 资料归档"); line("2. 确认版本")
            row["texts"] = [re.sub(r"^\d+\. ", "", t) for t in row["texts"]]
            row["selectors"].update(ol=1, li=2)
        c.save()
    for row in cases: row["sha256"] = digest(folder / row["source"])
    (folder / "expectations.json").write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8")
    return cases


def check_output(row, path, corpus, result):
    from bs4 import BeautifulSoup
    md = path.read_text(encoding="utf-8")
    soup = BeautifulSoup(MarkdownIt("commonmark", {"html": True}).enable("table").render(md), "html.parser")
    actual = norm(soup.get_text())
    checks = []
    def check(name, good, detail=""):
        checks.append(dict(name=name, passed=bool(good), detail=detail if not good else ""))
    check("全文逐字符一致（仅忽略空白与Unicode规范化）", actual == norm("".join(row["texts"])),
          "实际正文与预先定义的全文不一致，可能有丢字、多字或顺序变化")
    cursor = 0
    for i, text in enumerate(row["texts"], 1):
        expected = norm(text)
        at = actual.find(expected, cursor)
        check("正文片段%03d及顺序" % i, at >= 0, text)
        if at >= 0: cursor = at + len(expected)
    for selector, count in row["selectors"].items():
        actual_count = len(soup.select(selector))
        check("结构 " + selector, actual_count == count, "expected=%d actual=%d" % (count, actual_count))
    rows = [[norm(cell.get_text()) for cell in tr.find_all(["th", "td"], recursive=False)] for tr in soup.select("tr")]
    check("表格全部单元格和行列顺序", rows == [[norm(c) for c in r] for r in row["table_rows"]], repr(rows))
    for link in row["links"]: check("链接目标 " + link, any(a.get("href") == link for a in soup.select("a")))
    for code in row["codes"]: check("代码逐字与缩进", any(c.get_text().rstrip("\n") == code for c in soup.select("pre code")))
    imgs = soup.select("img")
    check("图片数量", len(imgs) == len(row["pictures"]), str(len(imgs)))
    for i, pic in enumerate(row["pictures"]):
        good = False
        if i < len(imgs):
            img_path = path.parent / imgs[i].get("src", "")
            if img_path.is_file():
                with Image.open(img_path) as got, Image.open(corpus / pic) as expected:
                    good = got.size == expected.size and got.convert("RGB").tobytes() == expected.convert("RGB").tobytes()
        check("图片%d像素一致" % (i+1), good)
    check("源文件未修改", digest(corpus / row["source"]) == row["sha256"])
    return checks


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "artifacts/conversion-qualification")
    parser.add_argument("--generate", action="store_true")
    parser.add_argument("--package", type=Path, default=ROOT / "plugins/packages/mdreader.document-convert-1.3.0.zip")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    corpus = args.out / "corpus"
    cases = make_corpus(corpus) if args.generate else json.loads((corpus / "expectations.json").read_text(encoding="utf-8"))
    if len(cases) != 20: raise ValueError("验收必须恰好包含 20 份源文档")
    run = args.out / ("run-" + time.strftime("%Y%m%d-%H%M%S"))
    run.mkdir()
    temporary = tempfile.TemporaryDirectory(prefix="mdreader-qualification-")
    ws = core.Workspace(temporary.name)
    records = []
    try:
        ws.plugins.install(str(args.package.resolve()))
        ws.plugins.enable("mdreader.document-convert")
        for row in cases:
            target = run / (row["id"] + ".md")
            began = time.monotonic()
            try:
                if digest(corpus / row["source"]) != row["sha256"]: raise ValueError("源样本哈希不匹配")
                result = convert_file(ws.plugins, "mdreader.document-convert:" + row["format"] + "-md", corpus / row["source"], target)
                checks = check_output(row, target, corpus, result)
                warnings = result["warnings"]
            except Exception as exc:
                checks = [dict(name="完成转换和检查", passed=False, detail=str(exc))]
                warnings = []
            record = dict(id=row["id"], name=row["name"], format=row["format"], complexity=row["complexity"],
                          passed=all(c["passed"] for c in checks), checks=checks, warnings=warnings,
                          seconds=round(time.monotonic()-began, 3), source_sha256=row["sha256"],
                          output_sha256=digest(target) if target.exists() else None)
            records.append(record)
            print(row["id"], "PASS" if record["passed"] else "FAIL", len(checks), flush=True)
            for check in checks:
                if not check["passed"]: print("  ", check["name"], check["detail"][:220], flush=True)
    finally:
        ws.plugins.shutdown()
        temporary.cleanup()
    count = sum(r["passed"] for r in records)
    check_count = sum(len(r["checks"]) for r in records)
    passed_checks = sum(c["passed"] for r in records for c in r["checks"])
    report = dict(package=str(args.package), package_sha256=digest(args.package),
                  corpus_sha256=digest(corpus / "expectations.json"), documents=20, passed_documents=count,
                  document_pass_rate=count/20, checks=check_count, passed_checks=passed_checks,
                  check_pass_rate=passed_checks/check_count, qualified=count/20 > .99 and passed_checks/check_count > .99,
                  scope="20 generated regression documents (10 DOCX + 10 text-layer PDF); not a population accuracy estimate. OCR, legacy DOC and native math objects are outside this corpus.",
                  results=records)
    (run / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# Word PDF 转 Markdown 验收报告", "", "固定样本：10 份 DOCX + 10 份文字层 PDF。",
             "文档必须通过所有适用检查才算通过；文档通过率与检查通过率均须严格大于 99%。",
             "这是一组生成式回归样本的结果，不代表任意真实文件达到 99% 准确率。扫描件 OCR、旧 DOC、原生公式对象未纳入此次范围。", "",
             "- 文档：%d/20（%.2f%%）" % (count, count*5),
             "- 检查点：%d/%d（%.2f%%）" % (passed_checks, check_count, passed_checks/check_count*100),
             "- 结论：" + ("本样本集达标" if report["qualified"] else "不合格，未达到门槛"),
             "- 安装包 SHA256：`" + report["package_sha256"] + "`", "",
             "| 编号 | 内容 | 格式 | 难度 | 通过检查 | 结果 |", "|---|---|---|---|---|---|"]
    for r in records:
        lines.append("| %s | %s | %s | %s | %d/%d | %s |" % (r["id"],r["name"],r["format"],r["complexity"],sum(c["passed"] for c in r["checks"]),len(r["checks"]),"通过" if r["passed"] else "失败"))
    lines += ["", "完整断言、失败原因、警告、耗时和输入输出哈希见同目录 report.json。原文件及预先定义的预期结果在 ../corpus。"]
    (run / "REPORT.md").write_text("\n".join(lines)+"\n", encoding="utf-8")
    print(str(run / "REPORT.md"), flush=True)
    return 0 if report["qualified"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
